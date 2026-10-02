"""End-to-end pipeline tests against the local fixture site."""

import openpyxl
import pandas as pd

from seo_fanout.export.tables import FINAL_COLUMNS
from seo_fanout.models import Action, Coverage, Risk
from seo_fanout.pipeline import ResearchInput, ResearchPipeline

SHEETS = ["FINAL 100", "WEBSITE INVENTORY", "CANNIBALIZATION", "CONTENT GAPS", "TOPIC CLUSTERS", "SEARCH VOLUME",
          "SERP RESEARCH", "COMMUNITY RESEARCH", "COMPETITOR GAPS", "QUERY SOURCES", "RECOMMENDED ACTIONS"]


def _run(settings, base, serp=None, **kw):
    steps = []
    pipe = ResearchPipeline(settings, lambda i, label: steps.append(i))
    if serp is not None:
        pipe._init_serp = lambda: serp  # type: ignore[method-assign]
    try:
        result, paths = pipe.run(ResearchInput(seed="sunscreen for oily skin", base_url=base, country="India",
                                               language="English", num_queries=kw.get("n", 100),
                                               fresh_crawl=kw.get("fresh", False)))
    finally:
        pipe.close()
    return result, paths, steps


def test_pipeline_without_any_api_keys(fixture_site, settings):
    result, paths, steps = _run(settings, fixture_site)
    assert steps == list(range(1, 11))
    for key in ("xlsx", "csv", "cannibalization", "content_gap", "summary"):
        assert paths[key].exists()
    # Cannibalization of the two duplicate posts is found before any query is generated
    high = [p for p in result.pairs if p.risk == Risk.HIGH]
    assert any("best-sunscreen" in p.url_a + p.url_b and "top-sunscreens" in p.url_a + p.url_b for p in high)
    # Never fabricated: no volumes, no SERP rows
    assert all(c.metrics.search_volume is None for c in result.final)
    assert not (result.serp and result.serp.serps)
    # Honest shortfall instead of padding
    assert len(result.final) < 100
    assert result.stats["selection"]["shortfall_reason"]
    # The head query is routed to cannibalization review, not to a new page
    head = next(c for c in result.candidates if c.query == "best sunscreen for oily skin")
    assert head.coverage == Coverage.CANNIBALIZATION_RISK and head.action == Action.CANNIBALIZATION_REVIEW

    wb = openpyxl.load_workbook(paths["xlsx"])
    assert wb.sheetnames == SHEETS
    df = pd.read_csv(paths["csv"])
    assert list(df.columns[: len(FINAL_COLUMNS)]) == FINAL_COLUMNS
    assert df["recommended_action"].notna().all() and df["reasoning"].notna().all()
    summary = paths["summary"].read_text()
    assert "# SEO Fan-Out Research Report" in summary and "Cannibalization Summary" in summary


def test_pipeline_with_serp_and_competitors(fixture_site, competitor_site, settings, fake_serp_factory, tmp_path):
    from seo_fanout.db import Database

    serp = fake_serp_factory(Database(tmp_path / "serp.sqlite3"), competitor_site)
    result, paths, _ = _run(settings, fixture_site, serp=serp)
    assert result.serp.available and result.serp.paa
    assert result.competitors.available and any(p.fetched for p in result.competitors.pages)
    assert result.community.questions  # quora/reddit/forum results via SERP API
    origins = set().union(*(c.origins for c in result.candidates))
    assert {"paa", "related", "autocomplete", "competitors"} <= origins
    # A competitor-only question becomes a content gap
    greasy = [c for c in result.final if "greasy" in c.query or "oily" in c.query and c.query.startswith("why")]
    assert greasy and all(c.coverage in (Coverage.CONTENT_GAP, Coverage.PARTIALLY_COVERED) for c in greasy)
    assert result.internal_links
    df = pd.read_csv(paths["csv"])
    assert (df["source_paa"] == "YES").any() and (df["source_competitors"] == "YES").any()


def test_crawl_is_cached_between_runs(fixture_site, settings):
    first, _, _ = _run(settings, fixture_site, n=20)
    second, _, _ = _run(settings, fixture_site, n=20)
    assert not first.crawl_reused and second.crawl_reused
    assert len(first.pages) == len(second.pages)
    fresh, _, _ = _run(settings, fixture_site, n=20, fresh=True)
    assert not fresh.crawl_reused


class FakeLLM:
    """TEST-ONLY LLM provider returning canned JSON."""

    name, model, calls = "fake-llm", "fake-model", 0

    def complete_json(self, system, user, max_tokens=4000):
        from seo_fanout.providers.llm import LLMResponse

        if "conversational_questions" in user:
            data = {"conversational_questions": ["what spf should i use for oily skin in humid weather"],
                    "follow_up_questions": ["can i use sunscreen instead of moisturizer on oily skin"],
                    "comparison_questions": ["is gel or fluid sunscreen better for oily skin"],
                    "entities": [{"entity": "niacinamide", "relationship": "controls sebum"}]}
        else:
            data = {"queries": [{"query": "does tinted sunscreen work for oily skin", "dimension": "core"},
                                {"query": "sunscreen stick for oily skin touch ups", "dimension": "situational"}]}
        return data, LLMResponse(text="{}", provider=self.name, model=self.model)


class FakePlanner:
    """TEST-ONLY keyword planner returning canned metrics for a subset of keywords."""

    calls = 0

    def fetch(self, keywords, country, language):
        from seo_fanout.models import DataLabel, KeywordMetrics
        from seo_fanout.textutil import normalize_query

        out = {}
        for i, kw in enumerate(keywords):
            if i % 2 == 0:
                out[normalize_query(kw)] = KeywordMetrics(keyword=kw, search_volume=1000 // (i + 1), competition="HIGH",
                                                          competition_index=80, cpc_low=0.2, cpc_high=1.1,
                                                          source="google_ads", label=DataLabel.REAL)
        return out


def test_pipeline_with_llm_and_keyword_metrics(fixture_site, settings):
    pipe = ResearchPipeline(settings)
    pipe._init_llms = lambda: [FakeLLM()]  # type: ignore[method-assign]
    pipe._init_ads = lambda: (FakePlanner(), "")  # type: ignore[method-assign]
    try:
        result, paths = pipe.run(ResearchInput(seed="sunscreen for oily skin", base_url=fixture_site, country="India",
                                               language="English", num_queries=50))
    finally:
        pipe.close()
    llm_queries = [c for c in result.candidates if "llm" in c.origins]
    assert llm_queries and all(c.label.value in ("LLM_GENERATED", "REAL_DATA") for c in llm_queries)
    with_vol = [c for c in result.final if c.metrics and c.metrics.search_volume is not None]
    without_vol = [c for c in result.final if c.metrics and c.metrics.search_volume is None]
    assert with_vol and without_vol
    assert all("search_volume" in c.score_breakdown for c in with_vol)
    assert all("search_volume" not in c.score_breakdown for c in without_vol)
    df = pd.read_csv(paths["csv"])
    assert (df["source_llm"] == "YES").any()
    assert df["search_volume"].notna().any() and df["search_volume"].isna().any()
