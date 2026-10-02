"""XLSX workbook exporter (11 sheets, formatted)."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..results import ResearchResult
from .tables import all_sheets

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RISK_FILL = {"HIGH": "F8D7DA", "MEDIUM": "FFF3CD", "LOW": "E2EFDA"}


def _illegal_free(df: pd.DataFrame) -> pd.DataFrame:
    """openpyxl rejects control characters – strip them from string cells."""
    import re

    pattern = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
    return df.apply(lambda col: col.map(lambda v: pattern.sub("", v) if isinstance(v, str) else v))


def write_excel(result: ResearchResult, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    sheets = all_sheets(result)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            _illegal_free(df).to_excel(writer, sheet_name=name[:31], index=False)
            ws = writer.sheets[name[:31]]
            ws.freeze_panes = "B2"
            for cell in ws[1]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for idx, column in enumerate(df.columns, start=1):
                sample = [str(column)] + [str(v) for v in df[column].head(200).tolist()]
                width = min(max(len(s) for s in sample) + 2, 60)
                ws.column_dimensions[get_column_letter(idx)].width = max(width, 10)
            if len(df.columns):
                ws.auto_filter.ref = ws.dimensions
            # Highlight risk columns
            for col_name in ("risk", "cannibalization_risk"):
                if col_name in df.columns:
                    col_idx = list(df.columns).index(col_name) + 1
                    for row in range(2, len(df) + 2):
                        cell = ws.cell(row=row, column=col_idx)
                        color = RISK_FILL.get(str(cell.value))
                        if color:
                            cell.fill = PatternFill("solid", fgColor=color)
    return path
